import os
from tkinter import Frame, ttk, filedialog, messagebox, StringVar
from numpy import copy
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
import pandas as pd

class MergeFiles(Frame):

    def __init__(self, controller, menu_bar):
        super().__init__(controller, highlightbackground="gray24", highlightthickness=1, padx=5, pady=5)
        self.columnconfigure(0, weight=1)
        
        self.menu_bar = menu_bar
        
        self.count_id_file_name = StringVar()
        self.count_id_ws = {}
        self.count_id_wb = {}
        
        self.merge_file_names = StringVar()
        self.merge_file_paths = []

        self._init_layout()
        
    def _init_layout(self):
        ttk.Label(self, text="CountID File Name", anchor="center").grid(row=0, column=0, padx=5, pady=10, sticky="nesw")
        ttk.Label(self, textvariable=self.count_id_file_name, anchor="center").grid(row=1, column=0, padx=5, pady=5, sticky="nesw")
        ttk.Button(self, text="Select File", command=self._upload_action).grid(row=2, column=0, padx=5, pady=5, sticky="nesw")
        ttk.Button(self, text="Remove CountID Column", command=self._remove_count_id).grid(row=3, column=0, padx=5, pady=5, sticky="nesw")
        ttk.Label(self, text="File Names", anchor="center").grid(row=4, column=0, padx=5, pady=10, sticky="nesw")
        ttk.Label(self, textvariable=self.merge_file_names, anchor="center").grid(row=5, column=0, padx=5, pady=5, sticky="nesw")
        ttk.Button(self, text="Select Files To Merge", command=self._upload_action_multiple).grid(row=6, column=0, padx=5, pady=5, sticky="nesw")
        ttk.Button(self, text="Merge", command=self._merge_work_books).grid(row=7, column=0, padx=5, pady=5, sticky="nesw")
        
    def _upload_action(self):
        # Uploading the orignial Excel file or orinigal file with <Count_ID> column
        file_path = filedialog.askopenfilename(title="Select Excel File", filetypes=[("Excel files", "*.xlsx")])
        if not file_path:
            return
        
        self.count_id_file_name.set(file_path.split("/")[-1])
        self.count_id_wb = load_workbook(file_path, data_only=True)
        self.count_id_ws = self.count_id_wb.active
        
    def _upload_action_multiple(self):
        # Upload multiple files that are selecting when mergine
        files = filedialog.askopenfilenames(filetypes=[("Excel files", "*.xlsx")])
        file_paths = ""

        # Adds all the file paths to a variable
        for path in files:
            file_paths += path.split("/")[-1] + "\n"
            self.merge_file_paths.append(path)

        self.merge_file_names.set(file_paths)
        
    def _remove_count_id(self):
        if self.menu_bar.output_dir_path:
            if self.count_id_wb:
                df = self._create_data_frame()
                if "Count_ID" in df.columns:
                    self.count_id_ws.delete_cols(1)
                    try:
                        file_path = os.path.join(self.menu_bar.output_dir_path, self.count_id_file_name.get())
                        self.count_id_wb.save(file_path)
                        if os.path.exists(file_path):
                            messagebox.showinfo("Success", f"Count_ID column removed. File saved as {self.count_id_file_name.get()}")
                    except Exception as e:
                        messagebox.showerror("Save Failed", f"An error occurred while saving the file:\n{str(e)}")
                else:
                    messagebox.showerror("Missing CountID column", "This file does not contain the Count_ID column.")
            else:
                messagebox.showerror("Missing CountID file", "No CountID file selected. Count_ID column must be included.")
        else:
            messagebox.showerror("Missing Output file", "No Output file selected.")
            
    def _create_data_frame(self):
        # Creates a dataframe so data can be grouped and column names can be read quickly
        
        if not self.count_id_ws:
            return
        data = self.count_id_ws.values
        columns = next(data)
        return pd.DataFrame(data, columns=columns)
            
    def _merge_work_books(self):
        if self.menu_bar.output_dir_path:
            if self.merge_file_paths:
                merge_ws = []
                for path in self.merge_file_paths:
                    wb = load_workbook(path)
                    ws = wb.active
                    merge_ws.append(ws)
                
                df = self._create_data_frame()

                if self.count_id_wb and "Count_ID" in df.columns:
                
                        update_cells = {}

                        for ws_idx, ws in enumerate(merge_ws):
                            for row_idx, row in enumerate(ws.rows):
                                update_cells[row[0].value] = {"ws_idx": ws_idx, "row_idx": row_idx + 1}


                        for col in merge_ws[0].columns:
                            max_length = max(len(str(cell.value)) for cell in col)

                            self.count_id_ws.column_dimensions[get_column_letter(col[0].column)].width = max_length

                        for row in self.count_id_ws.rows:
                            if row[0].value in update_cells:
                                for cell_new, cell in zip(merge_ws[update_cells[row[0].value]["ws_idx"]][update_cells[row[0].value]["row_idx"]], row):
                                        if cell_new.has_style:
                                            cell.font = copy(cell_new.font)
                                            cell.border = copy(cell_new.border)
                                            cell.fill = copy(cell_new.fill)
                                            cell.number_format = copy(cell_new.number_format)
                                            cell.protection = copy(cell_new.protection)
                                            cell.alignment = copy(cell_new.alignment)
                                        cell.value = cell_new.value
                        try:
                            file_path = os.path.join(self.menu_bar.output_dir_path, self.count_id_file_name.get())
                            self.count_id_wb.save(file_path)
                            if os.path.exists(file_path):
                                messagebox.showinfo("Success", "Files merged successfully. File saved as " + self.count_id_file_name.get())
                            else:
                                raise Exception("File not saved correctly.")
                        except Exception as e:
                            messagebox.showerror("Save Failed", f"An error occurred while saving the file:\n{str(e)}")
                        finally:
                            # Reset the variables after merging. Allows for a new merge to be done without issues.
                            self.count_id_file_name.set("")
                            self.count_id_wb = {}
                            self.count_id_ws = {}
                            self.merge_file_paths = []
                else:
                    messagebox.showerror("Missing CountID file", "No CountID file selected. <Count_ID> column must be included.")
            else:
                messagebox.showerror("Missing Merge files", "No files to merge selected.")
        else:
            messagebox.showerror("Missing Output file", "No Output file selected.")
