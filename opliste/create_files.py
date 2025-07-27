from tkinter import Frame, ttk, filedialog, messagebox, StringVar
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
import pandas as pd
import os


class CreateFiles(Frame):

    def __init__(self, controller, menu_bar):
        super().__init__(controller, highlightbackground="gray24", highlightthickness=1, padx=5, pady=5)
        self.columnconfigure(0, weight=1)
        
        self.menu_bar = menu_bar
        # Variable to hold the column name used for grouping
        self.column_name = StringVar()
        
        
        self.original_file_name = StringVar()
        self.original_ws = {}
        self.original_wb = {}
        
        self._init_layout()
        
    def _init_layout(self):
        ttk.Label(self, text="File Name", anchor="center").grid(row=0, column=0, padx=5, pady=10, sticky="nesw")
        ttk.Label(self, textvariable=self.original_file_name, anchor="center").grid(row=1, column=0, padx=5, pady=5, sticky="nesw")
        ttk.Button(self, text="Select File", command=self._upload_action).grid(row=2, column=0, padx=5, pady=5, sticky="nesw")
        ttk.Label(self, text="Column Name", anchor="center").grid(row=3, column=0, padx=5, pady=5, sticky="nesw")
        ttk.Entry(self, textvariable=self.column_name).grid(row=4, column=0, padx=5,pady=5, sticky="nesw")
        ttk.Button(self, text="Create Groups", command=self._create_work_books).grid(row=5, column=0, padx=5, pady=5, sticky="nesw")
        
    def _upload_action(self):
        # Uploading the orignial Excel file or orinigal file with <Count_ID> column
        file_path = filedialog.askopenfilename(title="Select Excel File", filetypes=[("Excel files", "*.xlsx")])
        if not file_path:
            return

        self.original_file_name.set(file_path.split("/")[-1])
        self.original_wb = load_workbook(file_path, data_only=True)
        self.original_ws = self.original_wb.active
        
    def _create_data_frame(self):
        # Creates a dataframe so data can be grouped and column names can be read quickly
        if not self.original_ws:
            return
        data = self.original_ws.values
        columns = next(data)
        
        return pd.DataFrame(data, columns=columns)
            
    def _create_work_books(self):
        df = self._create_data_frame()
        if self.menu_bar.output_dir_path:
            # Create a 'copy' of the original file with the column 'Count_ID'
            # Count_ID used to keep track of row loction when splitting and merging
            if "Count_ID" not in df.columns:
                df.insert(0, "Count_ID", range(1, len(df) + 1))
                df.to_excel(os.path.join(self.menu_bar.output_dir_path, "Astro_" + self.original_file_name.get()), index=False)
            # column_name is used to determine what column to group the data on
            if self.column_name.get() not in df.columns:
                if "Zuordnung" in df.columns:
                    self.column_name.set("Zuordnung")
                else:
                    self.column_name.set(df.columns[0])
            # Grouping data based on column name
            for groupValue, groupDF in df.groupby(self.column_name.get()):
                file_name = self.menu_bar.output_dir_path + "/" + str(groupValue) + ".xlsx"
                groupDF.to_excel(file_name, index=False)
                wb = load_workbook(file_name)
                ws = wb.active
                # Stlying the workbook
                for col in ws.columns:
                    # Set the column width to the maximum length of the data in that column. If cell is empty, it will be 0
                    max_length = max(len(str(cell.value)) if cell.value else 0 for cell in col)
                    if col[0].value == "Count_ID":
                        ws.column_dimensions[get_column_letter(col[0].column)].hidden= True
                    ws.column_dimensions[get_column_letter(col[0].column)].width = max_length
                wb.save(file_name)
            messagebox.showinfo("Success", "Mulitple excel files created. All files saved to " + self.menu_bar.output_dir_path)
        else:
            messagebox.showerror("Missing Output file", "No Output file selected.")